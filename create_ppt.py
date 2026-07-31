from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()

# Slide 1: Title
slide_layout_title = prs.slide_layouts[0]
slide1 = prs.slides.add_slide(slide_layout_title)
title1 = slide1.shapes.title
subtitle1 = slide1.placeholders[1]
title1.text = "PeopleRisk AI"
subtitle1.text = "Enterprise HR Intelligence\n\nPredictive analytics and real-time insights for organizational attrition."

# Slide 2: Problem & Vision
slide_layout_bullet = prs.slide_layouts[1]
slide2 = prs.slides.add_slide(slide_layout_bullet)
title2 = slide2.shapes.title
title2.text = "Problem & Vision"
body_shape2 = slide2.shapes.placeholders[1]
tf2 = body_shape2.text_frame
tf2.text = "The Challenge:"
p = tf2.add_paragraph()
p.text = "Unpredictable employee attrition leads to significant organizational costs and loss of institutional knowledge."
p.level = 1
p2 = tf2.add_paragraph()
p2.text = "The Solution:"
p2.level = 0
p3 = tf2.add_paragraph()
p3.text = "A data-driven platform that leverages live XGBoost models to predict flight risk and provide actionable insights."
p3.level = 1

# Slide 3: Key Features
slide3 = prs.slides.add_slide(slide_layout_bullet)
title3 = slide3.shapes.title
title3.text = "Key Features"
body_shape3 = slide3.shapes.placeholders[1]
tf3 = body_shape3.text_frame
tf3.text = "Interactive Dashboard:"
p4 = tf3.add_paragraph()
p4.text = "Real-time demographics, risk composition, and organizational heatmaps."
p4.level = 1
p5 = tf3.add_paragraph()
p5.text = "AI Assistant:"
p5.level = 0
p6 = tf3.add_paragraph()
p6.text = "Natural language querying for deep-dives into attrition drivers and risk tiers."
p6.level = 1
p7 = tf3.add_paragraph()
p7.text = "Automated Workflows:"
p7.level = 0
p8 = tf3.add_paragraph()
p8.text = "One-click email summaries, Slack alerts, and DOCX/PDF reporting."
p8.level = 1

# Slide 4: Technology Stack
slide4 = prs.slides.add_slide(slide_layout_bullet)
title4 = slide4.shapes.title
title4.text = "Technology Stack"
body_shape4 = slide4.shapes.placeholders[1]
tf4 = body_shape4.text_frame

tf4.text = "Streamlit:"
p9 = tf4.add_paragraph()
p9.text = "Rapid UI framework for building the interactive web dashboard."
p9.level = 1

p10 = tf4.add_paragraph()
p10.text = "Plotly:"
p10.level = 0
p11 = tf4.add_paragraph()
p11.text = "Graphing library for rendering dynamic charts and heatmaps."
p11.level = 1

p12 = tf4.add_paragraph()
p12.text = "XGBoost:"
p12.level = 0
p13 = tf4.add_paragraph()
p13.text = "Machine learning engine used for predicting employee flight risk."
p13.level = 1

p14 = tf4.add_paragraph()
p14.text = "SQLite:"
p14.level = 0
p15 = tf4.add_paragraph()
p15.text = "Lightweight database for storing HR records securely."
p15.level = 1

p16 = tf4.add_paragraph()
p16.text = "LLMs (Large Language Models):"
p16.level = 0
p17 = tf4.add_paragraph()
p17.text = "Powers the conversational AI assistant for natural language queries."
p17.level = 1

p18 = tf4.add_paragraph()
p18.text = "Pandas & NumPy:"
p18.level = 0
p19 = tf4.add_paragraph()
p19.text = "Core data manipulation libraries for real-time statistical aggregation."
p19.level = 1

prs.save("PeopleRiskAI_Introduction.pptx")
print("Presentation generated successfully!")
